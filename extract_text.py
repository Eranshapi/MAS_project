import os
from pptx import Presentation
import fitz  # PyMuPDF for PDFs
from docx import Document
import pytesseract
from pdf2image import convert_from_path
import logging
import tempfile
import sys
import shutil

# Configure logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)

# Configure Tesseract path for Windows
if os.name == 'nt':  # Windows
    pytesseract.pytesseract.tesseract_cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe'

def check_dependencies():
    """Check if all required dependencies are installed and accessible."""
    missing_deps = []
    
    # Check Tesseract
    try:
        tesseract_path = shutil.which('tesseract')
        if not tesseract_path:
            missing_deps.append("Tesseract OCR")
        else:
            # Configure Tesseract path
            pytesseract.pytesseract.tesseract_cmd = tesseract_path
    except Exception as e:
        logger.error(f"Error checking Tesseract: {str(e)}")
        missing_deps.append("Tesseract OCR")

    # Check Poppler
    try:
        poppler_path = shutil.which('pdftoppm')
        if not poppler_path:
            missing_deps.append("Poppler")
    except Exception as e:
        logger.error(f"Error checking Poppler: {str(e)}")
        missing_deps.append("Poppler")

    if missing_deps:
        error_msg = "Missing required dependencies:\n"
        for dep in missing_deps:
            error_msg += f"- {dep}\n"
        error_msg += "\nPlease install the missing dependencies:\n"
        error_msg += "1. Tesseract OCR: https://github.com/UB-Mannheim/tesseract/wiki\n"
        error_msg += "2. Poppler: https://github.com/oschwartz10612/poppler-windows/releases/\n"
        error_msg += "\nAfter installation, make sure to add them to your system PATH."
        logger.error(error_msg)
        return False
    return True

def extract_text_from_docx(file_path):
    doc = Document(file_path)
    text = []
    for paragraph in doc.paragraphs:
        text.append(paragraph.text)
    return "\n".join(text)

def extract_text_from_txt(file_path):
    with open(file_path, "r", encoding="utf-8") as f:
        return f.read()

def extract_text_from_pptx(file_path):
    prs = Presentation(file_path)
    text = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text.append(shape.text)
    return "\n".join(text)

def extract_text_from_pdf(file_path):
    """
    Extract text from PDF files, handling both regular and scanned PDFs.
    Uses PyMuPDF for regular PDFs and OCR for scanned PDFs.
    Supports Hebrew text recognition.
    """
    if not check_dependencies():
        return ""

    try:
        # First try regular PDF extraction
        doc = fitz.open(file_path)
        text = []
        
        for page_num in range(len(doc)):
            page = doc[page_num]
            
            # Try regular text extraction first
            page_text = page.get_text("text")
            
            # If no text is found, the page might be scanned
            if not page_text.strip():
                logger.info(f"Page {page_num + 1} appears to be scanned, using OCR...")
                
                try:
                    # Convert PDF page to image
                    with tempfile.TemporaryDirectory() as temp_dir:
                        # Convert PDF page to image
                        images = convert_from_path(
                            file_path,
                            first_page=page_num + 1,
                            last_page=page_num + 1,
                            dpi=300,  # Higher DPI for better OCR results
                            output_folder=temp_dir
                        )
                        
                        if images:
                            # Perform OCR on the image with Hebrew language support
                            page_text = pytesseract.image_to_string(
                                images[0],
                                lang='heb+eng',  # Support both Hebrew and English
                                config='--psm 1'  # Automatic page segmentation with OSD
                            )
                except Exception as e:
                    logger.error(f"OCR failed for page {page_num + 1}: {str(e)}")
                    page_text = ""
            
            text.append(page_text)
        
        doc.close()
        return "\n".join(text)
    
    except Exception as e:
        logger.error(f"Error extracting text from PDF {file_path}: {str(e)}")
        return ""

def extract_text_from_folder(folder_path):
    """Extracts text from all .txt, .pptx, .pdf, and .docx files in a folder and its subfolders."""
    all_texts = []
    
    for root, dirs, files in os.walk(folder_path):
        for file in files:
            file_path = os.path.join(root, file)
            if file.endswith(".txt"):
                all_texts.append(extract_text_from_txt(file_path))
            elif file.endswith(".pptx"):
                all_texts.append(extract_text_from_pptx(file_path))
            elif file.endswith(".pdf"):
                all_texts.append(extract_text_from_pdf(file_path))
            elif file.endswith(".docx"):
                all_texts.append(extract_text_from_docx(file_path))
    return all_texts
